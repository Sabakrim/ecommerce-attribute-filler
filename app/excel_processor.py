import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import logging
import re
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Any
from app.sku_matcher import normalize_sku, is_sku_match, clean_sku_key

logger = logging.getLogger("AttributeFiller.ExcelProcessor")

SKU_HEADER_ALIASES = {
    "sku", "item sku", "product sku", "sku #", "sku number", "sku id", "sku code", "sku name",
    "item #", "item code", "item no", "item number", "part #", "part no", "part number", "part id",
    "code", "product code", "product id", "product no", "product number",
    "model", "model #", "model no", "model number", "mpn", "mfr part #",
    "article no", "article number", "supplier sku", "seller sku", "vendor sku", "style #", "style no", "id", "name", "product"
}

class ExcelProcessingError(Exception):
    """Custom exception for Excel processing errors."""
    pass

def find_header_and_sku_column(sheet: openpyxl.worksheet.worksheet.Worksheet) -> Tuple[int, int, List[Tuple[int, str]]]:
    for row_idx in range(1, min(35, sheet.max_row + 1)):
        row_cells = [sheet.cell(row=row_idx, column=c) for c in range(1, sheet.max_column + 1)]
        row_values = [str(c.value).strip() if c.value is not None else "" for c in row_cells]
        
        for col_idx, val in enumerate(row_values, start=1):
            if val.lower() in SKU_HEADER_ALIASES:
                headers = []
                for c_i, h_val in enumerate(row_values, start=1):
                    if h_val:
                        headers.append((c_i, h_val))
                return row_idx, col_idx, headers

    for r in range(1, min(15, sheet.max_row + 1)):
        headers = []
        for c in range(1, sheet.max_column + 1):
            v = sheet.cell(row=r, column=c).value
            if v is not None and str(v).strip():
                headers.append((c, str(v).strip()))
        if headers:
            return r, headers[0][0], headers

    raise ExcelProcessingError("Could not detect header row in the Excel template.")

def create_standalone_specs_excel(output_path: Path, target_sku: str, extracted_specs: Dict[str, str]) -> None:
    """Creates a clean 2-column Excel file containing all extracted specifications for copy-pasting."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted Specifications"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(["Target SKU", target_sku])
    ws.append([])
    ws.append(["Attribute Name", "Attribute Value"])

    ws.cell(row=3, column=1).fill = header_fill
    ws.cell(row=3, column=1).font = header_font
    ws.cell(row=3, column=2).fill = header_fill
    ws.cell(row=3, column=2).font = header_font

    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 50

    r_idx = 4
    for k, v in extracted_specs.items():
        ws.append([k, v])
        c1 = ws.cell(row=r_idx, column=1)
        c2 = ws.cell(row=r_idx, column=2)
        c1.border = thin_border
        c2.border = thin_border
        r_idx += 1

    wb.save(output_path)

def process_excel_template(
    excel_path: Path,
    output_path: Path,
    target_sku: str,
    mapped_attributes: Dict[str, Tuple[str, float]],
    extracted_specs: Optional[Dict[str, str]] = None,
    target_worksheet_name: Optional[str] = None,
    overwrite_existing: bool = False
) -> Dict[str, Any]:
    if extracted_specs is None:
        extracted_specs = {}
    """
    Fills matching cells in uploaded Excel template AND appends an 'Extracted_Specs' sheet
    with ALL extracted key-value pairs so no data is ever lost.
    """
    if not target_sku or not target_sku.strip():
        raise ExcelProcessingError("Requested SKU cannot be empty.")

    target_sku_clean = target_sku.strip()
    target_sku_alphanumeric = clean_sku_key(target_sku_clean)

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        wb_evaluated = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception:
        # Fallback to standalone Excel if template file opening fails
        create_standalone_specs_excel(output_path, target_sku_clean, extracted_specs)
        return {
            "worksheet_name": "Extracted_Specs",
            "target_row": 1,
            "attributes_found": len(extracted_specs),
            "attributes_left_blank": 0,
            "filled_attributes": extracted_specs,
            "blank_attributes": []
        }

    sheet_names = wb.sheetnames
    matching_worksheets = []

    for s_name in sheet_names:
        ws = wb[s_name]
        ws_eval = wb_evaluated[s_name]

        try:
            h_row, sku_col, headers = find_header_and_sku_column(ws)
            sku_rows = []
            for r_idx in range(h_row + 1, ws.max_row + 1):
                v1 = ws.cell(row=r_idx, column=sku_col).value
                v2 = ws_eval.cell(row=r_idx, column=sku_col).value
                if (is_sku_match(target_sku_clean, v1) or 
                    is_sku_match(target_sku_clean, v2) or 
                    (v1 and target_sku_alphanumeric in clean_sku_key(v1)) or 
                    (v2 and target_sku_alphanumeric in clean_sku_key(v2))):
                    sku_rows.append(r_idx)

            if sku_rows:
                matching_worksheets.append((s_name, ws, h_row, sku_col, headers, sku_rows))
                continue
        except Exception:
            pass

        # 2D Cell search fallback
        found_cell = None
        for r_idx in range(1, ws.max_row + 1):
            for c_idx in range(1, ws.max_column + 1):
                v1 = ws.cell(row=r_idx, column=c_idx).value
                v2 = ws_eval.cell(row=r_idx, column=c_idx).value
                s1 = clean_sku_key(v1) if v1 is not None else ""
                s2 = clean_sku_key(v2) if v2 is not None else ""
                if (is_sku_match(target_sku_clean, v1) or 
                    is_sku_match(target_sku_clean, v2) or 
                    (s1 and target_sku_alphanumeric in s1) or 
                    (s2 and target_sku_alphanumeric in s2)):
                    found_cell = (r_idx, c_idx)
                    break
            if found_cell:
                break

        if found_cell:
            t_row, s_col = found_cell
            h_row = 1
            for search_r in range(t_row - 1, 0, -1):
                row_vals = [str(ws.cell(row=search_r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
                if any(row_vals):
                    h_row = search_r
                    break

            headers = []
            for c in range(1, ws.max_column + 1):
                h_val = ws.cell(row=h_row, column=c).value
                if h_val is not None and str(h_val).strip():
                    headers.append((c, str(h_val).strip()))

            matching_worksheets.append((s_name, ws, h_row, s_col, headers, [t_row]))

    wb_evaluated.close()

    if matching_worksheets:
        selected_ws_data = matching_worksheets[0]
        s_name, ws, h_row, sku_col, headers, sku_rows = selected_ws_data
        target_row = sku_rows[0]

        attributes_found = 0
        attributes_left_blank = 0
        filled_dict: Dict[str, str] = {}
        blank_list: List[str] = []

        for c_idx, h_name in headers:
            if c_idx == sku_col:
                continue

            cell = ws.cell(row=target_row, column=c_idx)
            existing_val = cell.value
            is_blank = existing_val is None or str(existing_val).strip() == ""

            if is_blank or overwrite_existing:
                if h_name in mapped_attributes:
                    extracted_val, confidence = mapped_attributes[h_name]
                    if extracted_val and str(extracted_val).strip():
                        cell.value = str(extracted_val).strip()
                        filled_dict[h_name] = str(extracted_val).strip()
                        attributes_found += 1
                    else:
                        attributes_left_blank += 1
                        blank_list.append(h_name)
                else:
                    attributes_left_blank += 1
                    blank_list.append(h_name)
    else:
        s_name = sheet_names[0]
        target_row = 1
        attributes_found = 0
        attributes_left_blank = 0
        filled_dict = {}
        blank_list = []

    # ALWAYS append/create an 'Extracted_Specs' sheet with ALL extracted raw specs for easy copy-pasting
    if "Extracted_Specs" in wb.sheetnames:
        del wb["Extracted_Specs"]
    
    spec_ws = wb.create_sheet(title="Extracted_Specs")
    spec_ws.append(["Target SKU", target_sku_clean])
    spec_ws.append([])
    spec_ws.append(["Attribute Name", "Attribute Value"])

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    spec_ws.cell(row=3, column=1).fill = header_fill
    spec_ws.cell(row=3, column=1).font = header_font
    spec_ws.cell(row=3, column=2).fill = header_fill
    spec_ws.cell(row=3, column=2).font = header_font
    spec_ws.column_dimensions['A'].width = 32
    spec_ws.column_dimensions['B'].width = 50

    for k, v in extracted_specs.items():
        spec_ws.append([k, v])

    wb.save(output_path)

    return {
        "worksheet_name": s_name,
        "target_row": target_row,
        "attributes_found": attributes_found if matching_worksheets else len(extracted_specs),
        "attributes_left_blank": attributes_left_blank,
        "filled_attributes": filled_dict if matching_worksheets else extracted_specs,
        "blank_attributes": blank_list
    }
